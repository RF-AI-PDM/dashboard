"""Tests for excel_writer.py, always against the test_excel throwaway copy."""

from __future__ import annotations

import os

import openpyxl
import pandas as pd
import pytest

import data_loader as dl
import excel_writer as ew


def _clean_monitoring_df(data: dict) -> pd.DataFrame:
    """Drop the in-memory-only "(Asli)" columns, as app.py does before writing."""
    mon = data["monitoring"]
    clean_cols = [c for c in mon.columns if not c.endswith("(Asli)")]
    return mon[clean_cols].copy()


def test_write_monitoring_roundtrip(test_excel, test_backup_dir):
    data = dl.load_all(test_excel)
    df = _clean_monitoring_df(data)
    first_title = df.iloc[0]["Judul Pekerjaan"]

    df.at[df.index[0], "STATUS"] = "SELESAI"
    backup_path = ew.write_monitoring(test_excel, df, backup_dir=test_backup_dir)

    assert os.path.exists(backup_path)
    assert os.path.dirname(backup_path) == test_backup_dir

    reloaded = dl.load_all(test_excel)
    row = reloaded["monitoring"][reloaded["monitoring"]["Judul Pekerjaan"] == first_title].iloc[0]
    assert row["STATUS"] == "SELESAI"
    # Row count must survive an edit that doesn't add/remove rows.
    assert len(reloaded["monitoring"]) == len(data["monitoring"])


def test_write_monitoring_row_count_change(test_excel, test_backup_dir):
    """Deleting a row must not corrupt the TOTAL row or the remaining data."""
    data = dl.load_all(test_excel)
    df = _clean_monitoring_df(data)
    original_len = len(df)
    kept_title = df.iloc[1]["Judul Pekerjaan"]

    df = df.iloc[1:].reset_index(drop=True)
    ew.write_monitoring(test_excel, df, backup_dir=test_backup_dir)

    reloaded = dl.load_all(test_excel)
    assert len(reloaded["monitoring"]) == original_len - 1
    assert kept_title in reloaded["monitoring"]["Judul Pekerjaan"].values


def test_append_histori_creates_sheet_and_row(test_excel, test_backup_dir):
    entry = {
        "timestamp": "01-01-2026 10:00:00",
        "judul": "Pekerjaan Uji",
        "field": "STATUS",
        "lama": "PROSES",
        "baru": "SELESAI",
    }
    ew.append_histori(test_excel, [entry], backup_dir=test_backup_dir)

    wb = openpyxl.load_workbook(test_excel)
    assert "Histori" in wb.sheetnames
    ws = wb["Histori"]
    assert ws.cell(row=1, column=1).value == "Timestamp"
    last_row = [ws.cell(row=ws.max_row, column=c).value for c in range(1, 6)]
    assert last_row == ["01-01-2026 10:00:00", "Pekerjaan Uji", "STATUS", "PROSES", "SELESAI"]
    wb.close()


def test_append_histori_empty_entries_noop(test_excel, test_backup_dir):
    result = ew.append_histori(test_excel, [], backup_dir=test_backup_dir)
    assert result == ""
    assert not os.path.exists(test_backup_dir)


def test_append_progress_note(test_excel, test_backup_dir):
    ew.append_progress_note(test_excel, judul="Pekerjaan Uji", note="Progres 50%",
                             author="Tester", backup_dir=test_backup_dir)
    wb = openpyxl.load_workbook(test_excel)
    ws = wb["Histori"]
    last_row = [ws.cell(row=ws.max_row, column=c).value for c in range(1, 6)]
    assert last_row[1] == "Pekerjaan Uji"
    assert last_row[2] == "Catatan (Tester)"
    assert last_row[4] == "Progres 50%"
    wb.close()


def test_write_targets(test_excel, test_backup_dir):
    ew.write_targets(
        test_excel,
        pencapaian={"target_tahun": 30000000000},
        backup_dir=test_backup_dir,
    )
    reloaded = dl.load_all(test_excel)
    assert reloaded["pencapaian"]["target_tahun"] == 30000000000


def test_backup_excel_creates_file(test_excel, test_backup_dir):
    backup_path = ew.backup_excel(test_excel, backup_dir=test_backup_dir)
    assert os.path.exists(backup_path)
    assert os.path.getsize(backup_path) == os.path.getsize(test_excel)


def test_write_monitoring_locked_file_raises(test_excel, test_backup_dir, monkeypatch):
    """Simulate Excel holding the file open (PermissionError -> ExcelLocked)."""
    data = dl.load_all(test_excel)
    df = _clean_monitoring_df(data)

    real_load_workbook = openpyxl.load_workbook

    def _raise_permission_error(*args, **kwargs):
        raise PermissionError("file is locked")

    monkeypatch.setattr(ew.openpyxl, "load_workbook", _raise_permission_error)
    with pytest.raises(ew.ExcelLocked):
        ew.write_monitoring(test_excel, df, backup_dir=test_backup_dir)
