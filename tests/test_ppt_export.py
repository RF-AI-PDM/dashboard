"""Tests for ppt_export.py: build succeeds, stays editable, uses locale-safe numbers."""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

import data_loader as dl
import ppt_export


def _build(test_excel) -> Presentation:
    data = dl.load_all(test_excel)
    raw_bytes = ppt_export.build_pptx(data, data["monitoring"])
    assert isinstance(raw_bytes, (bytes, bytearray))
    assert len(raw_bytes) > 10_000
    return Presentation(io.BytesIO(raw_bytes))


def test_build_pptx_smoke(test_excel):
    prs = _build(test_excel)
    assert len(prs.slides) >= 10
    assert prs.slide_width == Inches(13.333)
    assert prs.slide_height == Inches(7.5)


def test_chart_numbers_use_indonesian_locale(test_excel):
    """Regression test: chart labels must not fall back to a comma thousands separator."""
    prs = _build(test_excel)
    found_chart = False
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_chart:
                continue
            found_chart = True
            plot = shape.chart.plots[0]
            if plot.has_data_labels:
                assert plot.data_labels.number_format == ppt_export.ID_NUM_FORMAT
    assert found_chart, "expected at least one native chart in the export"


def test_logos_present_on_title_and_header_slides(test_excel):
    prs = _build(test_excel)
    title_slide = prs.slides[0]
    picture_count = sum(1 for shape in title_slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert picture_count >= 2, "title slide should carry both the Danantara and PLN IPS logos"

    kpi_slide = prs.slides[1]
    picture_count = sum(1 for shape in kpi_slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert picture_count >= 2, "content slide header should carry both logos"


def test_build_pptx_empty_monitoring_df(test_excel):
    """An empty (fully filtered) DataFrame must not crash the export."""
    data = dl.load_all(test_excel)
    empty_df = data["monitoring"].iloc[0:0]
    raw_bytes = ppt_export.build_pptx(data, empty_df)
    assert len(raw_bytes) > 10_000
