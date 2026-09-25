"""PowerPoint export utility for Dashboard Monitoring Pekerjaan.

Generates a clean, 16:9 widescreen PowerPoint presentation (.pptx)
using native python-pptx shapes, charts, and styled tables.
Follows the PLN / IPS visual identity (navy blue and yellow).
All source code must be ASCII-only.
"""

from __future__ import annotations

from datetime import datetime
import io
import math
import os
from typing import Any, Dict, List, Optional
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_DATA_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from data_loader import fmt_pct, fmt_rp
import theme

# Colors from theme
COLOR_PRIMARY = RGBColor(*theme.hex_to_rgb(theme.PRIMARY))          # #0A3D8F
COLOR_PRIMARY_DARK = RGBColor(*theme.hex_to_rgb(theme.PRIMARY_DARK)) # #062A63
COLOR_SECONDARY = RGBColor(*theme.hex_to_rgb(theme.SECONDARY))      # #1E88E5
COLOR_ACCENT = RGBColor(*theme.hex_to_rgb(theme.ACCENT))            # #FDB913
COLOR_SUCCESS = RGBColor(*theme.hex_to_rgb(theme.SUCCESS))          # #2E7D32
COLOR_WARNING = RGBColor(*theme.hex_to_rgb(theme.WARNING))          # #F57C00
COLOR_DANGER = RGBColor(*theme.hex_to_rgb(theme.DANGER))            # #C62828
COLOR_BG = RGBColor(*theme.hex_to_rgb(theme.BG))                    # #F4F7FB
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_DARK_TEXT = RGBColor(26, 26, 26)
COLOR_MUTED = RGBColor(110, 120, 135)
COLOR_BORDER = RGBColor(200, 215, 230)

MONTH_NAMES_ID = [
    "", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember"
]


def _format_date_id(dt: Optional[datetime] = None) -> str:
    """Format datetime as Indonesian date string (e.g. '22 September 2026')."""
    if dt is None:
        dt = datetime.now()
    month_name = MONTH_NAMES_ID[dt.month] if 1 <= dt.month <= 12 else ""
    return f"{dt.day} {month_name} {dt.year}"


def _truncate_text(text: Any, max_len: int = 60) -> str:
    """Truncate long string gracefully with ellipsis."""
    if text is None or pd.isna(text):
        return "-"
    s = str(text).strip()
    if not s or s.lower() == "nan":
        return "-"
    if len(s) > max_len:
        return s[:max_len - 3] + "..."
    return s


def _format_chart_data_and_axes(
    chart: Any,
    num_format: str = "#,##0",
    label_size_pt: float = 8.0,
    axis_size_pt: float = 8.0,
    label_position: Optional[Any] = None,
) -> None:
    """Format chart data labels, axes, and embedded workbook to avoid scientific notation."""
    try:
        plot = chart.plots[0]
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = num_format
        dl.font.size = Pt(label_size_pt)
        dl.font.name = "Segoe UI"
        dl.font.bold = True
        if label_position is not None:
            dl.position = label_position
    except Exception:
        pass

    try:
        va = chart.value_axis
        va.has_major_gridlines = True
        va.tick_labels.number_format = num_format
        va.tick_labels.font.size = Pt(axis_size_pt)
        va.tick_labels.font.name = "Segoe UI"
    except Exception:
        pass

    try:
        ca = chart.category_axis
        ca.tick_labels.font.size = Pt(axis_size_pt)
        ca.tick_labels.font.name = "Segoe UI"
    except Exception:
        pass

    # Ensure embedded Excel cells also have explicit number formatting
    try:
        cw = getattr(chart.part, "chart_workbook", None)
        if cw and getattr(cw, "xlsx_part", None):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(cw.xlsx_part.blob))
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, (int, float)):
                            cell.number_format = num_format
            buf = io.BytesIO()
            wb.save(buf)
            cw.update_from_xlsx_blob(buf.getvalue())
    except Exception:
        pass



def _add_header(slide: Any, title_text: str, subtitle_text: str = "") -> None:
    """Add standard PLN / IPS header banner (navy bar with thin yellow strip)."""
    # Navy bar full width (13.333 in x 0.9 in)
    navy_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.9)
    )
    navy_bar.fill.solid()
    navy_bar.fill.fore_color.rgb = COLOR_PRIMARY
    navy_bar.line.fill.background()

    # Thin yellow strip under it
    strip = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.9), Inches(13.333), Inches(0.06)
    )
    strip.fill.solid()
    strip.fill.fore_color.rgb = COLOR_ACCENT
    strip.line.fill.background()

    # Title text box
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.06), Inches(10.5), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.font.name = "Segoe UI"

    if subtitle_text:
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(10)
        p2.font.color.rgb = RGBColor(220, 235, 255)
        p2.font.name = "Segoe UI"

    # Logo top-right inside title bar if exists
    if os.path.exists(theme.LOGO_PATH):
        try:
            slide.shapes.add_picture(theme.LOGO_PATH, Inches(11.8), Inches(0.12), height=Inches(0.66))
        except Exception:
            pass


def _add_footer(slide: Any, slide_num: int, dt_str: str) -> None:
    """Add standard footer with organization name, date, and slide number."""
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.15), Inches(11.733), Inches(0.28))
    tf = tb.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"{theme.ORG_NAME}  |  Tanggal Laporan: {dt_str}  |  Slide {slide_num}"
    p.font.size = Pt(8.5)
    p.font.color.rgb = COLOR_MUTED
    p.font.name = "Segoe UI"


def _add_title_slide(prs: Presentation, dt_str: str) -> None:
    """Slide 1: Title slide in PLN / IPS style."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Background navy rectangle
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_PRIMARY
    bg.line.fill.background()

    # Decorative yellow bottom strip
    strip = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.35), Inches(13.333), Inches(0.15)
    )
    strip.fill.solid()
    strip.fill.fore_color.rgb = COLOR_ACCENT
    strip.line.fill.background()

    # Inner container card
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.1), Inches(10.933), Inches(5.2)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_WHITE
    card.line.color.rgb = COLOR_ACCENT
    card.line.width = Pt(2.5)

    # Logo in card if exists
    if os.path.exists(theme.LOGO_PATH):
        try:
            slide.shapes.add_picture(theme.LOGO_PATH, Inches(1.8), Inches(1.5), height=Inches(0.9))
        except Exception:
            pass

    # Text frame
    tb = slide.shapes.add_textbox(Inches(1.8), Inches(2.5), Inches(9.733), Inches(3.4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # Org tag
    p0 = tf.paragraphs[0]
    p0.text = theme.ORG_NAME.upper()
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_SECONDARY
    p0.font.name = "Segoe UI"
    p0.space_after = Pt(10)

    # Title
    p1 = tf.add_paragraph()
    p1.text = theme.APP_TITLE
    p1.font.size = Pt(34)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_PRIMARY
    p1.font.name = "Segoe UI"
    p1.space_after = Pt(10)

    # Subtitle
    p2 = tf.add_paragraph()
    p2.text = f"Laporan Kinerja, Progres Pekerjaan, dan Realisasi Penagihan\nTanggal Laporan: {dt_str}"
    p2.font.size = Pt(14)
    p2.font.color.rgb = COLOR_MUTED
    p2.font.name = "Segoe UI"


def _add_kpi_slide(
    prs: Presentation,
    data: Dict[str, Any],
    monitoring_df: pd.DataFrame,
    slide_num: int,
    dt_str: str
) -> None:
    """Slide 2: KPI summary cards (2x4 grid)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(
        slide,
        "Ringkasan Indikator Utama (KPI)",
        "Ikhtisar pencapaian target anggaran dan status operasional pekerjaan"
    )
    _add_footer(slide, slide_num, dt_str)

    pencapaian = data.get("pencapaian", {})
    potensi_df = data.get("potensi", pd.DataFrame())

    target_tahun = pencapaian.get("target_tahun", 0.0)
    realisasi_unfiltered = (
        float(monitoring_df["Nilai Kontrak Setelah PPN (IPS)"].sum())
        if "Nilai Kontrak Setelah PPN (IPS)" in monitoring_df.columns
        else 0.0
    )

    pct_pencapaian = (realisasi_unfiltered / target_tahun) if target_tahun > 0 else 0.0
    deviasi = realisasi_unfiltered - target_tahun

    total_invoice = (
        float(monitoring_df["Nilai Invoice Setelah PPN"].sum())
        if "Nilai Invoice Setelah PPN" in monitoring_df.columns
        else 0.0
    )
    total_potensi = (
        float(potensi_df["POTENSI NILAI"].sum())
        if not potensi_df.empty and "POTENSI NILAI" in potensi_df.columns
        else 0.0
    )

    jml_pekerjaan = len(monitoring_df)
    jml_selesai = (
        int((monitoring_df["STATUS"] == "SELESAI").sum())
        if "STATUS" in monitoring_df.columns
        else 0
    )

    # (label, value, is_highlighted)
    kpis = [
        ("Target Tahun 2026", fmt_rp(target_tahun), False),
        ("Realisasi Kontrak (IPS)", fmt_rp(realisasi_unfiltered), False),
        ("% Pencapaian", fmt_pct(pct_pencapaian), True),
        ("Deviasi Target", fmt_rp(deviasi), False),
        ("Total Nilai Invoice", fmt_rp(total_invoice), False),
        ("Total Potensi Pendapatan", fmt_rp(total_potensi), False),
        ("Jumlah Pekerjaan", f"{jml_pekerjaan} Pekerjaan", False),
        ("Pekerjaan Selesai", f"{jml_selesai} dari {jml_pekerjaan} Selesai", False),
    ]

    left_start = Inches(0.8)
    top_start = Inches(1.3)
    card_width = Inches(2.7)
    card_height = Inches(2.55)
    gap_x = Inches(0.3)
    gap_y = Inches(0.3)

    for idx, (label, val, is_highlight) in enumerate(kpis):
        r = idx // 4
        c = idx % 4
        x = left_start + c * (card_width + gap_x)
        y = top_start + r * (card_height + gap_y)

        rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_width, card_height)
        rect.fill.solid()

        if is_highlight:
            # Highlight card: navy fill, accent border
            rect.fill.fore_color.rgb = COLOR_PRIMARY
            rect.line.color.rgb = COLOR_ACCENT
            rect.line.width = Pt(2.5)
            val_color = COLOR_ACCENT
            label_color = COLOR_WHITE
        else:
            # Standard card: white fill, navy border
            rect.fill.fore_color.rgb = COLOR_WHITE
            rect.line.color.rgb = COLOR_PRIMARY
            rect.line.width = Pt(1.5)
            val_color = COLOR_PRIMARY
            label_color = COLOR_MUTED

        tb = slide.shapes.add_textbox(
            x + Inches(0.2), y + Inches(0.3), card_width - Inches(0.4), card_height - Inches(0.6)
        )
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = label.upper()
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = label_color
        p1.font.name = "Segoe UI"
        p1.space_after = Pt(14)

        p2 = tf.add_paragraph()
        p2.text = val
        p2.font.size = Pt(16 if len(val) > 13 else 22)
        p2.font.bold = True
        p2.font.color.rgb = val_color
        p2.font.name = "Segoe UI"


def _add_clustered_bar_slide(
    prs: Presentation,
    kpi_dict: Dict[str, Any],
    title: str,
    subtitle: str,
    slide_num: int,
    dt_str: str
) -> None:
    """Slide 3 & 4: Clustered Column Chart (Target = ACCENT, Realisasi = PRIMARY)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(slide, title, subtitle)
    _add_footer(slide, slide_num, dt_str)

    target_yr = kpi_dict.get("target_tahun", 0.0)
    real_yr = kpi_dict.get("realisasi_tahun", 0.0)
    target_s1 = kpi_dict.get("target_s1", 0.0)
    cap_s1 = kpi_dict.get("capaian_s1", 0.0)
    target_s2 = kpi_dict.get("target_s2", 0.0)
    cap_s2 = kpi_dict.get("capaian_s2", 0.0)

    chart_data = CategoryChartData()
    chart_data.categories = ["Tahun 2026", "Semester 1", "Semester 2"]
    chart_data.add_series("Target", (target_yr, target_s1, target_s2))
    chart_data.add_series("Realisasi", (real_yr, cap_s1, cap_s2))

    x = Inches(0.8)
    y = Inches(1.3)
    cx = Inches(8.5)
    cy = Inches(5.6)

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data
    ).chart

    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.plots[0].has_data_labels = True

    # Series colors: Target = ACCENT, Realisasi = PRIMARY
    plot = chart.plots[0]
    if len(plot.series) > 0:
        plot.series[0].format.fill.solid()
        plot.series[0].format.fill.fore_color.rgb = COLOR_ACCENT
    if len(plot.series) > 1:
        plot.series[1].format.fill.solid()
        plot.series[1].format.fill.fore_color.rgb = COLOR_PRIMARY

    # Format chart data labels, axes, and embedded Excel to avoid scientific notation
    _format_chart_data_and_axes(
        chart,
        num_format="#,##0",
        label_size_pt=8.5,
        axis_size_pt=8.0,
        label_position=XL_DATA_LABEL_POSITION.OUTSIDE_END
    )

    # Info card on the right
    rx = Inches(9.6)
    rcard = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rx, y, Inches(2.9), cy)
    rcard.fill.solid()
    rcard.fill.fore_color.rgb = COLOR_WHITE
    rcard.line.color.rgb = COLOR_PRIMARY
    rcard.line.width = Pt(1.5)

    tb = slide.shapes.add_textbox(rx + Inches(0.25), y + Inches(0.3), Inches(2.4), cy - Inches(0.6))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "RINGKASAN CAPAIAN"
    p0.font.bold = True
    p0.font.size = Pt(12)
    p0.font.color.rgb = COLOR_PRIMARY
    p0.space_after = Pt(14)

    metrics = [
        ("Tahun 2026", target_yr, real_yr, fmt_pct(kpi_dict.get("ratio_tahun", 0.0)), fmt_rp(kpi_dict.get("deviasi_tahun", 0.0))),
        ("Semester 1", target_s1, cap_s1, fmt_pct(kpi_dict.get("ratio_s1", 0.0)), fmt_rp(kpi_dict.get("deviasi_s1", 0.0))),
        ("Semester 2", target_s2, cap_s2, fmt_pct(kpi_dict.get("ratio_s2", 0.0)), fmt_rp(kpi_dict.get("deviasi_s2", 0.0))),
    ]

    for period, tgt, rls, pct_str, dev_str in metrics:
        p_per = tf.add_paragraph()
        p_per.text = period
        p_per.font.bold = True
        p_per.font.size = Pt(11)
        p_per.font.color.rgb = COLOR_SECONDARY

        p_val = tf.add_paragraph()
        p_val.text = f"Target: {fmt_rp(tgt)}\nRealisasi: {fmt_rp(rls)}\nCapaian: {pct_str} | Dev: {dev_str}"
        p_val.font.size = Pt(9.5)
        p_val.font.color.rgb = COLOR_DARK_TEXT
        p_val.space_after = Pt(10)


def _add_distribution_slide(
    prs: Presentation,
    monitoring_df: pd.DataFrame,
    slide_num: int,
    dt_str: str
) -> None:
    """Slide 5: Two native pie charts with themed slice colors."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(
        slide,
        "Distribusi Status Pekerjaan & Pembayaran",
        "Perbandingan status operasional dan progres penagihan/pembayaran"
    )
    _add_footer(slide, slide_num, dt_str)

    # Left: Status Pekerjaan
    status_counts = (
        monitoring_df["STATUS"].value_counts()
        if "STATUS" in monitoring_df.columns
        else pd.Series(dtype=int)
    )
    if not status_counts.empty:
        chart_data_1 = CategoryChartData()
        cats_1 = [_truncate_text(idx, 25) for idx in status_counts.index]
        chart_data_1.categories = cats_1
        chart_data_1.add_series("Status Pekerjaan", tuple(int(v) for v in status_counts.values))

        x1 = Inches(0.8)
        y1 = Inches(1.3)
        cx1 = Inches(5.6)
        cy1 = Inches(5.6)

        chart1 = slide.shapes.add_chart(
            XL_CHART_TYPE.PIE, x1, y1, cx1, cy1, chart_data_1
        ).chart
        chart1.has_legend = True
        chart1.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart1.legend.include_in_layout = False
        chart1.plots[0].has_data_labels = True
        chart1.plots[0].data_labels.number_format = "#,##0"

        # Apply status slice colors
        series1 = chart1.plots[0].series[0]
        for idx, pt in enumerate(series1.points):
            cat_name = str(status_counts.index[idx]).strip().upper()
            hex_col = theme.STATUS_COLORS.get(cat_name, theme.PALETTE[idx % len(theme.PALETTE)])
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor(*theme.hex_to_rgb(hex_col))

    # Right: Status Pembayaran
    pay_counts = (
        monitoring_df["STATUS PEMBAYARAN"].value_counts()
        if "STATUS PEMBAYARAN" in monitoring_df.columns
        else pd.Series(dtype=int)
    )
    if not pay_counts.empty:
        chart_data_2 = CategoryChartData()
        cats_2 = [_truncate_text(idx, 25) for idx in pay_counts.index]
        chart_data_2.categories = cats_2
        chart_data_2.add_series("Status Pembayaran", tuple(int(v) for v in pay_counts.values))

        x2 = Inches(6.9)
        y2 = Inches(1.3)
        cx2 = Inches(5.6)
        cy2 = Inches(5.6)

        chart2 = slide.shapes.add_chart(
            XL_CHART_TYPE.PIE, x2, y2, cx2, cy2, chart_data_2
        ).chart
        chart2.has_legend = True
        chart2.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart2.legend.include_in_layout = False
        chart2.plots[0].has_data_labels = True
        chart2.plots[0].data_labels.number_format = "#,##0"

        # Apply payment slice colors
        series2 = chart2.plots[0].series[0]
        for idx, pt in enumerate(series2.points):
            cat_name = str(pay_counts.index[idx]).strip().upper()
            hex_col = theme.PAYMENT_COLORS.get(cat_name, theme.PALETTE[idx % len(theme.PALETTE)])
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor(*theme.hex_to_rgb(hex_col))


def _add_top10_bar_slide(
    prs: Presentation,
    monitoring_df: pd.DataFrame,
    slide_num: int,
    dt_str: str
) -> None:
    """Slide 6: Horizontal bar chart of top 10 jobs by Nilai Kontrak Setelah PPN (IPS)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(
        slide,
        "Top 10 Pekerjaan berdasarkan Nilai Kontrak (IPS)",
        "Pekerjaan dengan alokasi nilai kontrak terbesar"
    )
    _add_footer(slide, slide_num, dt_str)

    col_val = "Nilai Kontrak Setelah PPN (IPS)"
    if col_val in monitoring_df.columns:
        valid_df = monitoring_df[monitoring_df[col_val] > 0].copy()
        valid_df = valid_df.sort_values(by=col_val, ascending=True).tail(10)
    else:
        valid_df = pd.DataFrame()

    if not valid_df.empty:
        chart_data = CategoryChartData()
        cats = [_truncate_text(j, 45) for j in valid_df["Judul Pekerjaan"]]
        chart_data.categories = cats
        chart_data.add_series("Nilai Kontrak Setelah PPN", tuple(float(v) for v in valid_df[col_val]))

        x = Inches(0.8)
        y = Inches(1.3)
        cx = Inches(11.733)
        cy = Inches(5.6)

        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED, x, y, cx, cy, chart_data
        ).chart
        chart.has_legend = False
        chart.plots[0].has_data_labels = True

        if len(chart.plots[0].series) > 0:
            chart.plots[0].series[0].format.fill.solid()
            chart.plots[0].series[0].format.fill.fore_color.rgb = COLOR_PRIMARY

        # Format chart data labels, axes, and embedded Excel to avoid scientific notation
        _format_chart_data_and_axes(
            chart,
            num_format="#,##0",
            label_size_pt=8.0,
            axis_size_pt=8.0,
            label_position=XL_DATA_LABEL_POSITION.OUTSIDE_END
        )


def _add_gantt_slide(
    prs: Presentation,
    gantt_png: bytes,
    slide_num: int,
    dt_str: str
) -> None:
    """Slide: Timeline / Gantt Schedule chart."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(
        slide,
        "Jadwal Pelaksanaan Pekerjaan (Timeline Gantt)",
        "Visualisasi jadwal pelaksanaan pekerjaan berdasarkan tanggal mulai dan berakhir"
    )
    _add_footer(slide, slide_num, dt_str)

    img_stream = io.BytesIO(gantt_png)
    slide.shapes.add_picture(img_stream, Inches(0.8), Inches(1.35), width=Inches(11.733))


def _add_kendala_slide(
    prs: Presentation,
    monitoring_df: pd.DataFrame,
    slide_num: int,
    dt_str: str
) -> None:
    """Slide: Kendala Utama Pekerjaan."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(
        slide,
        "Kendala & Hambatan Utama Pekerjaan",
        "Daftar pekerjaan aktif yang mengalami kendala operasional atau administratif"
    )
    _add_footer(slide, slide_num, dt_str)

    # Filter rows with non-empty Kendala and STATUS != SELESAI
    if "Kendala" in monitoring_df.columns and "STATUS" in monitoring_df.columns:
        mask = (
            monitoring_df["Kendala"].astype(str).str.strip() != ""
        ) & (
            monitoring_df["Kendala"].astype(str).str.strip().str.lower() != "nan"
        ) & (
            monitoring_df["STATUS"] != "SELESAI"
        )
        kendala_df = monitoring_df[mask].head(12)
    else:
        kendala_df = pd.DataFrame()

    if kendala_df.empty:
        # Card indicating no issues
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.5), Inches(10.333), Inches(2.5)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_BG
        card.line.color.rgb = COLOR_PRIMARY
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(Inches(2.0), Inches(3.2), Inches(9.333), Inches(1.2))
        p = tb.text_frame.paragraphs[0]
        p.text = "Tidak Ada Kendala Operasional Aktif Tercatat"
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = COLOR_PRIMARY
        p.font.name = "Segoe UI"
        p.alignment = PP_ALIGN.CENTER

        p2 = tb.text_frame.add_paragraph()
        p2.text = "Seluruh pekerjaan berjalan sesuai rencana atau telah selesai."
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_MUTED
        p2.font.name = "Segoe UI"
        p2.alignment = PP_ALIGN.CENTER
        return

    columns = ["NO", "Judul Pekerjaan", "STATUS", "Catatan Kendala"]
    col_widths = [0.7, 4.0, 1.4, 5.633]

    n_rows = len(kendala_df) + 1
    n_cols = len(columns)

    table_shape = slide.shapes.add_table(
        n_rows, n_cols, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.42 * n_rows)
    )
    table = table_shape.table
    _style_table(table, col_widths)

    # Header
    for c_i, col in enumerate(columns):
        _set_cell(
            table.cell(0, c_i),
            col,
            is_header=True,
            align=PP_ALIGN.CENTER if c_i in (0, 2) else PP_ALIGN.LEFT
        )

    # Rows
    for r_i, (_, row_data) in enumerate(kendala_df.iterrows()):
        t_row = r_i + 1
        is_even = (r_i % 2 == 1)
        _set_cell(table.cell(t_row, 0), str(row_data.get("NO", r_i + 1)), is_even=is_even, align=PP_ALIGN.CENTER)
        _set_cell(table.cell(t_row, 1), _truncate_text(row_data.get("Judul Pekerjaan"), 55), is_even=is_even)
        _set_cell(table.cell(t_row, 2), str(row_data.get("STATUS", "-")), is_even=is_even, align=PP_ALIGN.CENTER)
        _set_cell(table.cell(t_row, 3), _truncate_text(row_data.get("Kendala"), 120), is_even=is_even)


def _style_table(table: Any, col_widths: List[float]) -> None:
    """Apply column widths to table."""
    for idx, width in enumerate(col_widths):
        if idx < len(table.columns):
            table.columns[idx].width = Inches(width)


def _set_cell(
    cell: Any,
    text: str,
    is_header: bool = False,
    is_total: bool = False,
    is_even: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT
) -> None:
    """Set text, font, alignment, and background for a table cell."""
    cell.text = text
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for p in cell.text_frame.paragraphs:
        p.alignment = align
        p.font.name = "Segoe UI"
        if is_header:
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = COLOR_WHITE
        elif is_total:
            p.font.size = Pt(9)
            p.font.bold = True
            p.font.color.rgb = COLOR_PRIMARY
        else:
            p.font.size = Pt(8.5)
            p.font.color.rgb = COLOR_DARK_TEXT

    cell.fill.solid()
    if is_header:
        cell.fill.fore_color.rgb = COLOR_PRIMARY
    elif is_total:
        cell.fill.fore_color.rgb = RGBColor(230, 238, 250)
    elif is_even:
        cell.fill.fore_color.rgb = COLOR_BG
    else:
        cell.fill.fore_color.rgb = COLOR_WHITE


def _add_monitoring_table_slides(
    prs: Presentation,
    monitoring_df: pd.DataFrame,
    start_slide_num: int,
    dt_str: str
) -> int:
    """Paginated Monitoring table slides (10 rows per slide). Returns next slide number."""
    columns = [
        "NO", "Judul Pekerjaan", "User", "Proses Kontrak",
        "Nilai Kontrak Setelah PPN (IPS)", "Nama Vendor", "STATUS", "STATUS PEMBAYARAN"
    ]
    col_widths = [0.6, 3.2, 0.9, 1.4, 1.7, 1.6, 1.2, 1.1]

    rows_per_slide = 10
    total_rows = len(monitoring_df)
    total_pages = max(1, math.ceil(total_rows / rows_per_slide))

    curr_slide = start_slide_num
    for page in range(total_pages):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _add_header(
            slide,
            "Daftar Monitoring Pekerjaan",
            f"Halaman {page + 1} dari {total_pages} (Total {total_rows} pekerjaan terdaftar)"
        )
        _add_footer(slide, curr_slide, dt_str)
        curr_slide += 1

        start_idx = page * rows_per_slide
        end_idx = min(start_idx + rows_per_slide, total_rows)
        slice_df = monitoring_df.iloc[start_idx:end_idx]

        n_rows = len(slice_df) + 1
        n_cols = len(columns)

        table_shape = slide.shapes.add_table(
            n_rows, n_cols, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.45 * n_rows)
        )
        table = table_shape.table
        _style_table(table, col_widths)

        header_labels = [
            "NO", "Judul Pekerjaan", "User", "Proses Kontrak",
            "Nilai Kontrak (IPS)", "Nama Vendor", "Status", "Pembayaran"
        ]
        for col_idx, label in enumerate(header_labels):
            _set_cell(
                table.cell(0, col_idx),
                label,
                is_header=True,
                align=PP_ALIGN.CENTER if col_idx in (0, 2) else PP_ALIGN.LEFT
            )

        for row_i, (_, row_data) in enumerate(slice_df.iterrows()):
            table_row = row_i + 1
            is_even = (row_i % 2 == 1)
            for col_i, col_name in enumerate(columns):
                val = row_data.get(col_name, "-")
                if col_name == "Nilai Kontrak Setelah PPN (IPS)":
                    cell_text = fmt_rp(val)
                    align = PP_ALIGN.RIGHT
                elif col_name == "NO":
                    cell_text = str(val) if pd.notna(val) else "-"
                    align = PP_ALIGN.CENTER
                elif col_name == "Judul Pekerjaan":
                    cell_text = _truncate_text(val, 55)
                    align = PP_ALIGN.LEFT
                else:
                    cell_text = _truncate_text(val, 25)
                    align = PP_ALIGN.LEFT

                _set_cell(table.cell(table_row, col_i), cell_text, is_even=is_even, align=align)

    return curr_slide


def _add_monitoring_akumulatif_slide(
    prs: Presentation,
    data: Dict[str, Any],
    slide_num: int,
    dt_str: str
) -> None:
    """Table Monitoring Akumulatif + total row."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(
        slide,
        "Daftar Monitoring Akumulatif",
        "Pekerjaan akumulatif tahun jamak / kumulatif berjalan"
    )
    _add_footer(slide, slide_num, dt_str)

    df = data.get("monitoring_akumulatif", pd.DataFrame())
    columns = [
        "NO", "Judul Pekerjaan", "User", "Kontrak", "Nilai Kontrak Setelah PPN",
        "Nama Vendor", "Kontrak Vendor", "Nilai Kontrak Mitra Setelah PPN", "STATUS"
    ]
    col_widths = [0.5, 2.5, 0.7, 1.2, 1.6, 1.6, 1.4, 1.6, 0.9]

    total_val_ips = (
        float(df["Nilai Kontrak Setelah PPN"].sum())
        if "Nilai Kontrak Setelah PPN" in df.columns
        else 0.0
    )
    total_val_mitra = (
        float(df["Nilai Kontrak Mitra Setelah PPN"].sum())
        if "Nilai Kontrak Mitra Setelah PPN" in df.columns
        else 0.0
    )

    n_rows = len(df) + 2
    n_cols = len(columns)

    table_shape = slide.shapes.add_table(
        n_rows, n_cols, Inches(0.6), Inches(1.35), Inches(12.133), Inches(0.45 * n_rows)
    )
    table = table_shape.table
    _style_table(table, col_widths)

    for c_i, col in enumerate(columns):
        _set_cell(
            table.cell(0, c_i),
            col,
            is_header=True,
            align=PP_ALIGN.CENTER if c_i in (0, 2) else PP_ALIGN.LEFT
        )

    for r_i, (_, row) in enumerate(df.iterrows()):
        table_row = r_i + 1
        is_even = (r_i % 2 == 1)
        for c_i, col in enumerate(columns):
            val = row.get(col, "-")
            if "Nilai" in col:
                text = fmt_rp(val)
                align = PP_ALIGN.RIGHT
            elif col == "NO":
                text = str(val) if pd.notna(val) else "-"
                align = PP_ALIGN.CENTER
            else:
                text = _truncate_text(val, 40)
                align = PP_ALIGN.LEFT
            _set_cell(table.cell(table_row, c_i), text, is_even=is_even, align=align)

    tot_row = len(df) + 1
    _set_cell(table.cell(tot_row, 0), "", is_total=True)
    _set_cell(table.cell(tot_row, 1), "TOTAL", is_total=True, align=PP_ALIGN.LEFT)
    _set_cell(table.cell(tot_row, 2), "", is_total=True)
    _set_cell(table.cell(tot_row, 3), "", is_total=True)
    _set_cell(table.cell(tot_row, 4), fmt_rp(total_val_ips), is_total=True, align=PP_ALIGN.RIGHT)
    _set_cell(table.cell(tot_row, 5), "", is_total=True)
    _set_cell(table.cell(tot_row, 6), "", is_total=True)
    _set_cell(table.cell(tot_row, 7), fmt_rp(total_val_mitra), is_total=True, align=PP_ALIGN.RIGHT)
    _set_cell(table.cell(tot_row, 8), "", is_total=True)


def _add_potensi_slide(
    prs: Presentation,
    data: Dict[str, Any],
    slide_num: int,
    dt_str: str
) -> None:
    """Table Potensi Pendapatan + total row."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(
        slide,
        "Daftar Potensi Pendapatan 2025 - 2026",
        "Peluang pendapatan kontrak baru yang belum terkontrak"
    )
    _add_footer(slide, slide_num, dt_str)

    df = data.get("potensi", pd.DataFrame())
    columns = ["NO", "JUDUL", "POTENSI NILAI", "PERIODE"]
    col_widths = [0.8, 6.2, 3.0, 1.733]

    total_val = (
        float(df["POTENSI NILAI"].sum())
        if not df.empty and "POTENSI NILAI" in df.columns
        else 0.0
    )

    n_rows = len(df) + 2
    n_cols = len(columns)

    table_shape = slide.shapes.add_table(
        n_rows, n_cols, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.48 * n_rows)
    )
    table = table_shape.table
    _style_table(table, col_widths)

    for c_i, col in enumerate(columns):
        _set_cell(
            table.cell(0, c_i),
            col,
            is_header=True,
            align=PP_ALIGN.CENTER if c_i in (0, 3) else PP_ALIGN.LEFT
        )

    for r_i, (_, row) in enumerate(df.iterrows()):
        table_row = r_i + 1
        is_even = (r_i % 2 == 1)
        _set_cell(table.cell(table_row, 0), str(r_i + 1), is_even=is_even, align=PP_ALIGN.CENTER)
        _set_cell(table.cell(table_row, 1), _truncate_text(row.get("JUDUL", "-"), 60), is_even=is_even, align=PP_ALIGN.LEFT)
        _set_cell(table.cell(table_row, 2), fmt_rp(row.get("POTENSI NILAI", 0.0)), is_even=is_even, align=PP_ALIGN.RIGHT)
        _set_cell(table.cell(table_row, 3), str(row.get("PERIODE", "-")), is_even=is_even, align=PP_ALIGN.CENTER)

    tot_row = len(df) + 1
    _set_cell(table.cell(tot_row, 0), "", is_total=True)
    _set_cell(table.cell(tot_row, 1), "TOTAL POTENSI", is_total=True, align=PP_ALIGN.LEFT)
    _set_cell(table.cell(tot_row, 2), fmt_rp(total_val), is_total=True, align=PP_ALIGN.RIGHT)
    _set_cell(table.cell(tot_row, 3), "", is_total=True)


def _add_penagihan_slide(
    prs: Presentation,
    data: Dict[str, Any],
    slide_num: int,
    dt_str: str
) -> None:
    """Slide: Alur Penagihan (checklist sections)."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_header(
        slide,
        "Alur & Persyaratan Dokumen Penagihan",
        "Standar operasional prosedur pengajuan berkas invoice dan penagihan"
    )
    _add_footer(slide, slide_num, dt_str)

    sections = data.get("penagihan", [])
    if not sections:
        return

    left_start = Inches(0.8)
    top_pos = Inches(1.35)
    card_width = Inches(3.7)
    card_height = Inches(5.5)
    gap_x = Inches(0.3)

    for idx, sec in enumerate(sections[:3]):
        x = left_start + idx * (card_width + gap_x)

        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top_pos, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = COLOR_PRIMARY if idx == 0 else COLOR_BORDER
        card.line.width = Pt(1.8 if idx == 0 else 1.0)

        # Header Pill
        pill = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.2), top_pos + Inches(0.2), Inches(0.9), Inches(0.35)
        )
        pill.fill.solid()
        pill.fill.fore_color.rgb = COLOR_PRIMARY
        pill.line.fill.background()
        p_pill = pill.text_frame.paragraphs[0]
        p_pill.text = f"BAGIAN {sec.get('letter', '')}"
        p_pill.font.bold = True
        p_pill.font.size = Pt(9)
        p_pill.font.color.rgb = COLOR_WHITE
        p_pill.alignment = PP_ALIGN.CENTER

        # Title & Checklist
        tb = slide.shapes.add_textbox(
            x + Inches(0.2), top_pos + Inches(0.7), card_width - Inches(0.4), card_height - Inches(0.9)
        )
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = sec.get("title", "")
        p0.font.bold = True
        p0.font.size = Pt(11)
        p0.font.color.rgb = COLOR_PRIMARY
        p0.space_after = Pt(12)

        for item in sec.get("items", []):
            pi = tf.add_paragraph()
            pi.text = f"[  ]  {item}"
            pi.font.size = Pt(9.5)
            pi.font.color.rgb = COLOR_DARK_TEXT
            pi.space_after = Pt(8)


def build_pptx(
    data: Dict[str, Any],
    monitoring_df: pd.DataFrame,
    gantt_png: Optional[bytes] = None
) -> bytes:
    """Build complete presentation in PLN / IPS visual theme and return raw bytes."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    today_str = _format_date_id()
    slide_num = 1

    # 1. Title slide
    _add_title_slide(prs, today_str)
    slide_num += 1

    # 2. KPI slide
    _add_kpi_slide(prs, data, monitoring_df, slide_num, today_str)
    slide_num += 1

    # 3. Native clustered column chart: PENCAPAIAN 2026
    pencapaian = data.get("pencapaian", {})
    if pencapaian:
        _add_clustered_bar_slide(
            prs,
            pencapaian,
            "Target vs Realisasi Kinerja Tahun 2026",
            "Evaluasi realisasi pencapaian target anggaran tahun berjalan dan per semester",
            slide_num,
            today_str
        )
        slide_num += 1

    # 4. Native clustered column chart: Grafik Akumulatif
    akumulatif_kpi = data.get("akumulatif_kpi", {})
    if akumulatif_kpi:
        _add_clustered_bar_slide(
            prs,
            akumulatif_kpi,
            "Target vs Realisasi Akumulatif (Kumulatif 2025)",
            "Evaluasi akumulatif kinerja kontrak tahun jamak berjalan",
            slide_num,
            today_str
        )
        slide_num += 1

    # 5. Distribution pie charts (STATUS & STATUS PEMBAYARAN)
    _add_distribution_slide(prs, monitoring_df, slide_num, today_str)
    slide_num += 1

    # 6. Horizontal bar chart: Top 10 Nilai Kontrak
    _add_top10_bar_slide(prs, monitoring_df, slide_num, today_str)
    slide_num += 1

    # 7. Gantt timeline slide (if provided)
    if gantt_png:
        _add_gantt_slide(prs, gantt_png, slide_num, today_str)
        slide_num += 1

    # 8. Kendala Utama slide
    _add_kendala_slide(prs, monitoring_df, slide_num, today_str)
    slide_num += 1

    # 9. Monitoring table slides (paginated)
    slide_num = _add_monitoring_table_slides(prs, monitoring_df, slide_num, today_str)

    # 10. Monitoring Akumulatif table slide
    _add_monitoring_akumulatif_slide(prs, data, slide_num, today_str)
    slide_num += 1

    # 11. Potensi Pendapatan table slide
    _add_potensi_slide(prs, data, slide_num, today_str)
    slide_num += 1

    # 12. Alur Penagihan slide
    _add_penagihan_slide(prs, data, slide_num, today_str)

    # Save to buffer
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
