"""Excel writer module for Dashboard Monitoring Pekerjaan.

Writes back to data/Data.xlsx using openpyxl while preserving formulas and styles.
All source code must be ASCII-only.
"""

from __future__ import annotations

from copy import copy
from datetime import datetime
import io
import math
import os
import re
import shutil
from typing import Any, Dict, List, Optional
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd


class ExcelLocked(Exception):
    """Raised when the Excel file is locked or cannot be written."""
    pass


def backup_excel(path: str, backup_dir: Optional[str] = None) -> str:
    """Create a backup of the Excel file, keeping only the 20 newest backups."""
    if backup_dir is None:
        backup_dir = os.path.join(os.path.dirname(path), "backup")
    os.makedirs(backup_dir, exist_ok=True)

    stem = os.path.splitext(os.path.basename(path))[0]
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_filename = f"{stem}_{timestamp}.xlsx"
    backup_path = os.path.join(backup_dir, backup_filename)

    counter = 1
    while os.path.exists(backup_path):
        backup_path = os.path.join(backup_dir, f"{stem}_{timestamp}_{counter}.xlsx")
        counter += 1

    shutil.copy2(path, backup_path)

    # Prune old backups for this stem, keeping only newest 20
    matching_files = [
        os.path.join(backup_dir, f)
        for f in os.listdir(backup_dir)
        if f.startswith(f"{stem}_") and f.endswith(".xlsx") and os.path.isfile(os.path.join(backup_dir, f))
    ]
    matching_files.sort(key=lambda p: os.path.getmtime(p), reverse=True)
    if len(matching_files) > 20:
        for old_file in matching_files[20:]:
            try:
                os.remove(old_file)
            except OSError:
                pass

    return backup_path


def _find_total_row(ws: Any, col_letter: str, start_row: int) -> int:
    """Find row number where cell value in col_letter equals 'TOTAL'."""
    for r in range(start_row, ws.max_row + 1):
        val = ws[f"{col_letter}{r}"].value
        if val is not None and str(val).strip().upper() == "TOTAL":
            return r
    raise ValueError(f"Baris TOTAL tidak ditemukan di kolom {col_letter} mulai baris {start_row}")


def _ensure_slots(
    ws: Any,
    first_data_row: int,
    total_row: int,
    needed: int,
    merged_total_range: Optional[str] = None
) -> int:
    """Ensure sufficient row slots before the TOTAL row.

    If needed > (total_row - first_data_row), inserts rows before TOTAL.
    Unmerges and re-merges the total range, and copies styles from the last data row.
    Returns the new total_row.
    """
    slots = total_row - first_data_row
    if needed <= slots:
        return total_row

    amount = needed - slots
    actual_merged: Optional[tuple[int, int]] = None

    # Check and unmerge total range if present
    if merged_total_range:
        for rng in list(ws.merged_cells.ranges):
            if str(rng) == merged_total_range or (rng.min_row == total_row and rng.max_row == total_row):
                actual_merged = (rng.min_col, rng.max_col)
                ws.unmerge_cells(str(rng))
                break

    last_data_row = total_row - 1
    ws.insert_rows(total_row, amount)
    new_total_row = total_row + amount

    # Re-merge at the new total row
    if actual_merged:
        min_col, max_col = actual_merged
        min_letter = openpyxl.utils.get_column_letter(min_col)
        max_letter = openpyxl.utils.get_column_letter(max_col)
        ws.merge_cells(f"{min_letter}{new_total_row}:{max_letter}{new_total_row}")

    # Copy styles from last_data_row to each newly inserted row
    for r in range(total_row, new_total_row):
        for c in range(1, ws.max_column + 1):
            src_cell = ws.cell(row=last_data_row, column=c)
            tgt_cell = ws.cell(row=r, column=c)
            if src_cell.has_style:
                tgt_cell.font = copy(src_cell.font)
                tgt_cell.border = copy(src_cell.border)
                tgt_cell.fill = copy(src_cell.fill)
                tgt_cell.number_format = copy(src_cell.number_format)
                tgt_cell.alignment = copy(src_cell.alignment)

    return new_total_row


def _rewrite_total_formulas(
    ws: Any,
    total_row: int,
    first_data_row: int,
    numeric_col_letters: List[str]
) -> None:
    """Rewrite SUM formulas on the TOTAL row for specified numeric columns."""
    for col in numeric_col_letters:
        ws[f"{col}{total_row}"] = f"=SUM({col}{first_data_row}:{col}{total_row - 1})"


def _clean_num_val(v: Any) -> Optional[float]:
    """Convert value to float if valid numeric, else None."""
    if v is None or pd.isna(v):
        return None
    try:
        val = float(v)
        if math.isnan(val) or math.isinf(val):
            return None
        return val
    except (ValueError, TypeError):
        return None


def _clean_str_val(v: Any) -> Optional[str]:
    """Convert value to stripped string if non-empty, else None."""
    if v is None or pd.isna(v):
        return None
    s = str(v).strip()
    if not s or s.lower() == "nan":
        return None
    return s


# Monitoring column definition: (1-based Excel column index, column letter, column name, is_numeric)
MONITORING_COLS = [
    (2, "B", "NO", False),
    (3, "C", "Judul Pekerjaan", False),
    (4, "D", "Projeck", False),
    (5, "E", "User", False),
    (6, "F", "Proses Kontrak", False),
    (7, "G", "Eksekutor", False),
    (8, "H", "DMR/TOR/Spesifikasi", False),
    (9, "I", "LOI", False),
    (10, "J", "Nomor Informasi Harga Mitra", False),
    (11, "K", "IH Sebelum PPN", True),
    (12, "L", "IH Setelah PPN", True),
    (13, "M", "Undangan PL", False),
    (14, "N", "Kontrak IPS", False),
    (15, "O", "Nilai Kontrak Sebelum PPN (IPS)", True),
    (16, "P", "Nilai Kontrak Setelah PPN (IPS)", True),
    (17, "Q", "Mulai Pekerjaan", False),
    (18, "R", "Berakhir Pekerjaan", False),
    (19, "S", "Addendum Kontrak", False),
    (20, "T", "Ket Addendum", False),
    (21, "U", "BAST/TUG", False),
    (22, "V", "Invoice", False),
    (23, "W", "Nilai Invoice Sebelum PPN", True),
    (24, "X", "Nilai Invoice Setelah PPN", True),
    (25, "Y", "DENDA", True),
    (26, "Z", "PEMBAYARAN PIPS", True),
    (27, "AA", "Nama Vendor", False),
    (28, "AB", "Nomor Kontrak Vendor", False),
    (29, "AC", "Nilai Kontrak Mitra Sebelum PPN", True),
    (30, "AD", "Nilai Kontrak Mitra Setelah PPN", True),
    (31, "AE", "BA Vendor", False),
    (32, "AF", "Nilai BA Vendor Setelah PPN", True),
    (33, "AG", "DENDA2", True),
    (34, "AH", "PEMBAYARAN MITRA", True),
    (35, "AI", "MARGIN", True),
    (36, "AJ", "Kendala", False),
    (37, "AK", "Periode Penagihan", False),
    (38, "AL", "STATUS", False),
    (39, "AM", "STATUS PEMBAYARAN", False),
]

MONITORING_NUMERIC_TOTAL_COLS = ["K", "L", "O", "P", "W", "X", "Y", "Z", "AC", "AD", "AF"]


def write_monitoring(path: str, df: pd.DataFrame, backup_dir: Optional[str] = None) -> str:
    """Write DataFrame back to sheet 'Monitoring' in Excel file.

    Performs automatic backup, expands row slots if needed, clears previous data,
    writes updated rows with renumbered NO, updates TOTAL row SUM formulas, and
    updates formula references to Monitoring!P<row> across other sheets.
    """
    backup_path = backup_excel(path, backup_dir=backup_dir)

    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except PermissionError as e:
        raise ExcelLocked("Tutup file Excel terlebih dahulu") from e

    ws = wb["Monitoring"]
    first_data_row = 5
    old_total_row = _find_total_row(ws, "B", first_data_row)

    total_row = _ensure_slots(
        ws,
        first_data_row,
        old_total_row,
        len(df),
        merged_total_range=f"B{old_total_row}:H{old_total_row}"
    )

    # Ensure row 4 is template row NO=0 with empty title so it is kept as template
    ws.cell(row=4, column=2).value = 0
    ws.cell(row=4, column=3).value = None

    # Clear columns B..AM from row 5 to total_row - 1
    for r in range(first_data_row, total_row):
        for col_idx, col_letter, _, _ in MONITORING_COLS:
            ws.cell(row=r, column=col_idx).value = None

    # Write rows in df order
    for row_i, (_, row_data) in enumerate(df.iterrows()):
        curr_row = first_data_row + row_i
        for col_idx, col_letter, col_name, is_numeric in MONITORING_COLS:
            if col_name == "NO":
                ws.cell(row=curr_row, column=col_idx).value = row_i + 1
            elif is_numeric:
                val = _clean_num_val(row_data.get(col_name))
                ws.cell(row=curr_row, column=col_idx).value = val
            else:
                raw_val = row_data.get(col_name)
                s_val = _clean_str_val(raw_val)
                if col_name in ("STATUS", "STATUS PEMBAYARAN"):
                    if s_val == "BELUM ADA STATUS":
                        s_val = None
                ws.cell(row=curr_row, column=col_idx).value = s_val

    # Rewrite total formulas
    _rewrite_total_formulas(ws, total_row, first_data_row, MONITORING_NUMERIC_TOTAL_COLS)

    # If total_row moved or changed, update formula references in all sheets
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        for r in range(1, sheet.max_row + 1):
            for c in range(1, sheet.max_column + 1):
                cell_val = sheet.cell(row=r, column=c).value
                if isinstance(cell_val, str) and cell_val.startswith("="):
                    if "Monitoring!P" in cell_val:
                        new_val = re.sub(r"Monitoring!P\d+", f"Monitoring!P{total_row}", cell_val)
                        if new_val != cell_val:
                            sheet.cell(row=r, column=c).value = new_val

    try:
        wb.save(path)
    except PermissionError as e:
        raise ExcelLocked("Tutup file Excel terlebih dahulu") from e
    finally:
        wb.close()

    return backup_path


AKUMULATIF_COLS = [
    (2, "B", "NO", False),
    (3, "C", "Judul Pekerjaan", False),
    (4, "D", "User", False),
    (5, "E", "Kontrak", False),
    (6, "F", "Nilai Kontrak Setelah PPN", True),
    (7, "G", "Nama Vendor", False),
    (8, "H", "Kontrak Vendor", False),
    (9, "I", "Nilai Kontrak Mitra Setelah PPN", True),
    (10, "J", "STATUS", False),
]

AKUMULATIF_TOTAL_COLS = ["F", "G", "H", "I"]


def write_monitoring_akumulatif(path: str, df: pd.DataFrame, backup_dir: Optional[str] = None) -> str:
    """Write DataFrame back to sheet 'Monitoring Akumulatif '."""
    backup_path = backup_excel(path, backup_dir=backup_dir)

    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except PermissionError as e:
        raise ExcelLocked("Tutup file Excel terlebih dahulu") from e

    sheet_name = "Monitoring Akumulatif " if "Monitoring Akumulatif " in wb.sheetnames else "Monitoring Akumulatif"
    ws = wb[sheet_name]
    first_data_row = 2
    old_total_row = _find_total_row(ws, "B", first_data_row)

    total_row = _ensure_slots(
        ws,
        first_data_row,
        old_total_row,
        len(df),
        merged_total_range=f"B{old_total_row}:D{old_total_row}"
    )

    # Clear columns B..J from row 2 to total_row - 1
    for r in range(first_data_row, total_row):
        for col_idx, _, _, _ in AKUMULATIF_COLS:
            ws.cell(row=r, column=col_idx).value = None

    # Write rows
    for row_i, (_, row_data) in enumerate(df.iterrows()):
        curr_row = first_data_row + row_i
        for col_idx, _, col_name, is_numeric in AKUMULATIF_COLS:
            if col_name == "NO":
                ws.cell(row=curr_row, column=col_idx).value = row_i + 1
            elif is_numeric:
                ws.cell(row=curr_row, column=col_idx).value = _clean_num_val(row_data.get(col_name))
            else:
                s_val = _clean_str_val(row_data.get(col_name))
                if col_name == "STATUS" and s_val == "BELUM ADA STATUS":
                    s_val = None
                ws.cell(row=curr_row, column=col_idx).value = s_val

    # Rewrite totals for F, G, H, I
    _rewrite_total_formulas(ws, total_row, first_data_row, AKUMULATIF_TOTAL_COLS)

    # Update references in formulas e.g. 'Monitoring Akumulatif '!I8
    for s_name in wb.sheetnames:
        sheet = wb[s_name]
        for r in range(1, sheet.max_row + 1):
            for c in range(1, sheet.max_column + 1):
                cell_val = sheet.cell(row=r, column=c).value
                if isinstance(cell_val, str) and cell_val.startswith("="):
                    if "Monitoring Akumulatif" in cell_val and "!I" in cell_val:
                        new_val = re.sub(r"('?Monitoring Akumulatif\s*'?\s*!\s*I)\d+", rf"\g<1>{total_row}", cell_val)
                        if new_val != cell_val:
                            sheet.cell(row=r, column=c).value = new_val

    try:
        wb.save(path)
    except PermissionError as e:
        raise ExcelLocked("Tutup file Excel terlebih dahulu") from e
    finally:
        wb.close()

    return backup_path


def write_potensi(path: str, df: pd.DataFrame, backup_dir: Optional[str] = None) -> str:
    """Write DataFrame back to sheet 'POTENSI 2025 DAN 2026'."""
    backup_path = backup_excel(path, backup_dir=backup_dir)

    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except PermissionError as e:
        raise ExcelLocked("Tutup file Excel terlebih dahulu") from e

    sheet_name = "POTENSI 2025 DAN 2026"
    ws = wb[sheet_name]
    first_data_row = 5
    old_total_row = _find_total_row(ws, "C", first_data_row)

    total_row = _ensure_slots(ws, first_data_row, old_total_row, len(df), merged_total_range=None)

    # Clear columns C..E from row 5 to total_row - 1
    for r in range(first_data_row, total_row):
        ws.cell(row=r, column=3).value = None
        ws.cell(row=r, column=4).value = None
        ws.cell(row=r, column=5).value = None

    # Write data
    for row_i, (_, row_data) in enumerate(df.iterrows()):
        curr_row = first_data_row + row_i
        ws.cell(row=curr_row, column=3).value = _clean_str_val(row_data.get("JUDUL"))
        ws.cell(row=curr_row, column=4).value = _clean_num_val(row_data.get("POTENSI NILAI"))
        ws.cell(row=curr_row, column=5).value = _clean_str_val(row_data.get("PERIODE"))

    # Rewrite total formula for D
    ws[f"D{total_row}"] = f"=SUM(D{first_data_row}:D{total_row - 1})"

    try:
        wb.save(path)
    except PermissionError as e:
        raise ExcelLocked("Tutup file Excel terlebih dahulu") from e
    finally:
        wb.close()

    return backup_path


HISTORI_HEADERS = ["Timestamp", "Judul Pekerjaan", "Field", "Nilai Lama", "Nilai Baru"]


def append_histori(path: str, entries: List[Dict[str, Any]], backup_dir: Optional[str] = None) -> str:
    """Append change-log rows to sheet 'Histori', creating it if missing.

    Each entry: {"timestamp": str, "judul": str, "field": str, "lama": Any, "baru": Any}.
    Pure append-only log, no TOTAL row, so no row-shifting bookkeeping is needed.
    """
    if not entries:
        return ""

    backup_path = backup_excel(path, backup_dir=backup_dir)

    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except PermissionError as e:
        raise ExcelLocked("Tutup file Excel terlebih dahulu") from e

    if "Histori" in wb.sheetnames:
        ws = wb["Histori"]
    else:
        ws = wb.create_sheet("Histori")
        for col_idx, header in enumerate(HISTORI_HEADERS, start=1):
            ws.cell(row=1, column=col_idx).value = header

    next_row = ws.max_row + 1
    for entry in entries:
        ws.cell(row=next_row, column=1).value = _clean_str_val(entry.get("timestamp")) or "-"
        ws.cell(row=next_row, column=2).value = _clean_str_val(entry.get("judul")) or "-"
        ws.cell(row=next_row, column=3).value = _clean_str_val(entry.get("field")) or "-"
        ws.cell(row=next_row, column=4).value = _clean_str_val(entry.get("lama")) or "-"
        ws.cell(row=next_row, column=5).value = _clean_str_val(entry.get("baru")) or "-"
        next_row += 1

    try:
        wb.save(path)
    except PermissionError as e:
        raise ExcelLocked("Tutup file Excel terlebih dahulu") from e
    finally:
        wb.close()

    return backup_path


def append_progress_note(
    path: str,
    judul: str,
    note: str,
    author: str = "-",
    backup_dir: Optional[str] = None
) -> str:
    """Append a manual progress note for a job to sheet 'Histori'."""
    ts = datetime.now().strftime("%d-%m-%Y %H:%M:%S")
    clean_author = str(author).strip() if author else "-"
    field_label = f"Catatan ({clean_author})" if clean_author and clean_author != "-" else "Catatan Progres"
    entry = {
        "timestamp": ts,
        "judul": str(judul).strip(),
        "field": field_label,
        "lama": "-",
        "baru": str(note).strip(),
    }
    return append_histori(path, [entry], backup_dir=backup_dir)


def write_targets(
    path: str,
    pencapaian: Optional[Dict[str, Any]] = None,
    akumulatif: Optional[Dict[str, Any]] = None,
    backup_dir: Optional[str] = None
) -> str:
    """Update target constants in sheet 'PENCAPAIAN' and 'Grafik akumulatif'."""
    backup_path = backup_excel(path, backup_dir=backup_dir)

    try:
        wb = openpyxl.load_workbook(path, data_only=False)
    except PermissionError as e:
        raise ExcelLocked("Tutup file Excel terlebih dahulu") from e

    # Update PENCAPAIAN
    if pencapaian and "PENCAPAIAN" in wb.sheetnames:
        ws_p = wb["PENCAPAIAN"]
        if "target_tahun" in pencapaian:
            ws_p["C6"] = _clean_num_val(pencapaian["target_tahun"])
        if "target_s1" in pencapaian:
            ws_p["E6"] = _clean_num_val(pencapaian["target_s1"])
        if "target_s2" in pencapaian:
            ws_p["G6"] = _clean_num_val(pencapaian["target_s2"])

    # Update Grafik akumulatif
    if akumulatif and "Grafik akumulatif" in wb.sheetnames:
        ws_g = wb["Grafik akumulatif"]
        if "target_tahun" in akumulatif:
            ws_g["C6"] = _clean_num_val(akumulatif["target_tahun"])
        if "target_s1" in akumulatif:
            ws_g["E6"] = _clean_num_val(akumulatif["target_s1"])
        if "capaian_s1" in akumulatif:
            ws_g["F6"] = _clean_num_val(akumulatif["capaian_s1"])
        if "target_s2" in akumulatif:
            ws_g["G6"] = _clean_num_val(akumulatif["target_s2"])

    try:
        wb.save(path)
    except PermissionError as e:
        raise ExcelLocked("Tutup file Excel terlebih dahulu") from e
    finally:
        wb.close()

    return backup_path


def _fmt_rp_str(x: Any) -> str:
    """Format number as Indonesian Rupiah string (e.g. 'Rp 16.618.722.076')."""
    if x is None or pd.isna(x):
        return "-"
    try:
        val = float(x)
        if math.isnan(val) or math.isinf(val):
            return "-"
        val_int = int(round(val))
        formatted = f"{val_int:,}".replace(",", ".")
        return f"Rp {formatted}"
    except (ValueError, TypeError):
        return "-"


def build_job_sheet_html(job_row: Dict[str, Any], hist_df: pd.DataFrame) -> str:
    """Generate a clean, print-ready HTML Job Sheet (Lembar Kendali Pekerjaan)."""
    no_val = str(job_row.get("NO", "-"))
    judul_val = str(job_row.get("Judul Pekerjaan", "-"))
    projeck_val = str(job_row.get("Projeck", "-"))
    user_val = str(job_row.get("User", "-"))
    proses_val = str(job_row.get("Proses Kontrak", "-"))
    eksekutor_val = str(job_row.get("Eksekutor", "-"))
    kontrak_val = str(job_row.get("Kontrak IPS", "-"))
    vendor_val = str(job_row.get("Nama Vendor", "-"))
    status_val = str(job_row.get("STATUS", "-"))
    bayar_val = str(job_row.get("STATUS PEMBAYARAN", "-"))
    mulai_val = str(job_row.get("Mulai Pekerjaan", "-"))
    berakhir_val = str(job_row.get("Berakhir Pekerjaan", "-"))
    kendala_val = str(job_row.get("Kendala", "")).strip()

    nilai_ips = _fmt_rp_str(job_row.get("Nilai Kontrak Setelah PPN (IPS)"))
    nilai_seb = _fmt_rp_str(job_row.get("Nilai Kontrak Sebelum PPN (IPS)"))
    nilai_mitra = _fmt_rp_str(job_row.get("Nilai Kontrak Mitra Setelah PPN"))
    invoice_val = _fmt_rp_str(job_row.get("Nilai Invoice Setelah PPN"))
    margin_val = _fmt_rp_str(job_row.get("MARGIN"))

    hist_rows_html = []
    if hist_df is not None and not hist_df.empty:
        filtered_hist = hist_df[hist_df["Judul Pekerjaan"] == judul_val] if "Judul Pekerjaan" in hist_df.columns else hist_df
        if not filtered_hist.empty:
            sorted_hist = filtered_hist.sort_values("Timestamp", ascending=False)
            for idx, (_, h) in enumerate(sorted_hist.iterrows(), 1):
                ts = str(h.get("Timestamp", "-"))
                fld = str(h.get("Field", "-"))
                lama = str(h.get("Nilai Lama", "-"))
                baru = str(h.get("Nilai Baru", "-"))
                hist_rows_html.append(f"""
                <tr>
                    <td style="text-align: center; padding: 6px 8px; border: 1px solid #CBD5E1;">{idx}</td>
                    <td style="padding: 6px 8px; border: 1px solid #CBD5E1; white-space: nowrap;">{ts}</td>
                    <td style="padding: 6px 8px; border: 1px solid #CBD5E1; font-weight: 600;">{fld}</td>
                    <td style="padding: 6px 8px; border: 1px solid #CBD5E1; color: #64748B;">{lama}</td>
                    <td style="padding: 6px 8px; border: 1px solid #CBD5E1;">{baru}</td>
                </tr>
                """)

    if not hist_rows_html:
        hist_rows_html.append("""
        <tr>
            <td colspan="5" style="text-align: center; padding: 12px; color: #64748B; border: 1px solid #CBD5E1;">
                Belum ada riwayat histori / catatan progres yang tersimpan untuk pekerjaan ini.
            </td>
        </tr>
        """)

    hist_content = "".join(hist_rows_html)
    kendala_section = f"""
    <div style="background-color: #FEF3C7; border: 1px solid #F59E0B; border-radius: 6px; padding: 10px 14px; margin-bottom: 16px; font-size: 13px;">
        <strong style="color: #92400E;">Catatan / Kendala Terdaftar:</strong> {kendala_val}
    </div>
    """ if kendala_val else ""

    print_date = datetime.now().strftime("%d-%m-%Y %H:%M:%S")

    html = f"""<!DOCTYPE html>
<html lang="id">
<head>
    <meta charset="UTF-8">
    <title>Lembar Kendali Pekerjaan - {no_val}</title>
    <style>
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
            color: #1E293B;
            background-color: #F1F5F9;
            margin: 0;
            padding: 20px;
        }}
        .sheet-container {{
            max-width: 860px;
            margin: 0 auto;
            background: #FFFFFF;
            border: 1px solid #CBD5E1;
            border-radius: 8px;
            padding: 30px;
            box-shadow: 0 4px 6px -1px rgba(0,0,0,0.05);
        }}
        .header-bar {{
            display: flex;
            justify-content: space-between;
            align-items: center;
            border-bottom: 3px solid #0A3D8F;
            padding-bottom: 12px;
            margin-bottom: 20px;
        }}
        .header-title {{
            font-size: 20px;
            font-weight: 700;
            color: #0A3D8F;
            margin: 0;
        }}
        .header-subtitle {{
            font-size: 13px;
            color: #FDB913;
            font-weight: 600;
            margin: 3px 0 0 0;
            text-transform: uppercase;
            letter-spacing: 0.5px;
        }}
        .grid-info {{
            display: grid;
            grid-template-columns: 1fr 1fr;
            gap: 10px 24px;
            margin-bottom: 20px;
            font-size: 13px;
        }}
        .info-row {{
            display: flex;
            border-bottom: 1px solid #F1F5F9;
            padding-bottom: 4px;
        }}
        .info-label {{
            width: 150px;
            color: #64748B;
            font-weight: 600;
        }}
        .info-val {{
            flex: 1;
            font-weight: 600;
            color: #0F172A;
        }}
        .kpi-boxes {{
            display: grid;
            grid-template-columns: repeat(4, 1fr);
            gap: 12px;
            margin-bottom: 20px;
        }}
        .kpi-card {{
            background: #F8FAFC;
            border: 1px solid #E2E8F0;
            border-left: 4px solid #0A3D8F;
            border-radius: 6px;
            padding: 10px 12px;
        }}
        .kpi-label {{
            font-size: 11px;
            font-weight: 600;
            color: #64748B;
            text-transform: uppercase;
        }}
        .kpi-value {{
            font-size: 14px;
            font-weight: 700;
            color: #0A3D8F;
            margin-top: 4px;
        }}
        .section-title {{
            font-size: 14px;
            font-weight: 700;
            color: #0A3D8F;
            border-bottom: 2px solid #E2E8F0;
            padding-bottom: 4px;
            margin-top: 24px;
            margin-bottom: 12px;
            text-transform: uppercase;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 12px;
            margin-bottom: 24px;
        }}
        th {{
            background-color: #0A3D8F;
            color: #FFFFFF;
            padding: 8px;
            text-align: left;
            font-weight: 600;
            border: 1px solid #0A3D8F;
        }}
        .signature-block {{
            display: grid;
            grid-template-columns: 1fr 1fr;
            margin-top: 36px;
            text-align: center;
            font-size: 13px;
        }}
        .sign-space {{
            height: 60px;
        }}
        .no-print {{
            margin-bottom: 16px;
            text-align: right;
        }}
        .btn-print {{
            background: #0A3D8F;
            color: #FFFFFF;
            border: none;
            padding: 8px 18px;
            border-radius: 6px;
            font-weight: 600;
            cursor: pointer;
        }}
        @media print {{
            body {{ background: #FFFFFF; padding: 0; }}
            .no-print {{ display: none; }}
            .sheet-container {{ border: none; box-shadow: none; padding: 0; max-width: 100%; }}
            @page {{ size: A4 portrait; margin: 12mm 15mm; }}
        }}
    </style>
</head>
<body>
    <div class="sheet-container">
        <div class="no-print">
            <button class="btn-print" onclick="window.print()">Cetak / Simpan PDF</button>
        </div>
        <div class="header-bar">
            <div>
                <h1 class="header-title">LEMBAR KENDALI & MONITORING PEKERJAAN</h1>
                <div class="header-subtitle">PT PLN INDONESIA POWER SERVICES</div>
            </div>
            <div style="text-align: right; font-size: 12px; color: #64748B;">
                <div>No. Ref: <strong>#{no_val}</strong></div>
                <div>Tgl Cetak: {print_date}</div>
            </div>
        </div>

        <div class="grid-info">
            <div class="info-row"><span class="info-label">Judul Pekerjaan</span><span class="info-val">{judul_val}</span></div>
            <div class="info-row"><span class="info-label">Status Pekerjaan</span><span class="info-val" style="color: #0A3D8F;">{status_val}</span></div>
            <div class="info-row"><span class="info-label">User / Bidang</span><span class="info-val">{user_val}</span></div>
            <div class="info-row"><span class="info-label">Status Pembayaran</span><span class="info-val">{bayar_val}</span></div>
            <div class="info-row"><span class="info-label">Eksekutor</span><span class="info-val">{eksekutor_val}</span></div>
            <div class="info-row"><span class="info-label">Nama Vendor</span><span class="info-val">{vendor_val}</span></div>
            <div class="info-row"><span class="info-label">Nomor Kontrak IPS</span><span class="info-val">{kontrak_val}</span></div>
            <div class="info-row"><span class="info-label">Jadwal Mulai</span><span class="info-val">{mulai_val}</span></div>
            <div class="info-row"><span class="info-label">Proses Kontrak</span><span class="info-val">{proses_val}</span></div>
            <div class="info-row"><span class="info-label">Target Selesai</span><span class="info-val">{berakhir_val}</span></div>
        </div>

        <div class="kpi-boxes">
            <div class="kpi-card">
                <div class="kpi-label">Nilai Kontrak IPS (PPN)</div>
                <div class="kpi-value">{nilai_ips}</div>
            </div>
            <div class="kpi-card">
                <div class="kpi-label">Nilai Kontrak Mitra (PPN)</div>
                <div class="kpi-value">{nilai_mitra}</div>
            </div>
            <div class="kpi-card">
                <div class="kpi-label">Estimasi Margin</div>
                <div class="kpi-value">{margin_val}</div>
            </div>
            <div class="kpi-card">
                <div class="kpi-label">Nilai Invoice (PPN)</div>
                <div class="kpi-value">{invoice_val}</div>
            </div>
        </div>

        {kendala_section}

        <div class="section-title">Riwayat Histori & Log Aktivitas Pekerjaan</div>
        <table>
            <thead>
                <tr>
                    <th style="width: 35px; text-align: center;">No</th>
                    <th style="width: 130px;">Waktu</th>
                    <th style="width: 150px;">Aktivitas / Field</th>
                    <th style="width: 140px;">Nilai Lama / Pencatat</th>
                    <th>Keterangan / Nilai Baru</th>
                </tr>
            </thead>
            <tbody>
                {hist_content}
            </tbody>
        </table>

        <div class="signature-block">
            <div>
                <div>Disiapkan Oleh:</div>
                <div class="sign-space"></div>
                <div style="font-weight: 700;">( {user_val if user_val != '-' else 'PIC Pekerjaan'} )</div>
                <div style="font-size: 11px; color: #64748B;">Pengawas / Pelaksana</div>
            </div>
            <div>
                <div>Disetujui / Diverifikasi Oleh:</div>
                <div class="sign-space"></div>
                <div style="font-weight: 700;">( ........................................ )</div>
                <div style="font-size: 11px; color: #64748B;">Manager / Team Leader</div>
            </div>
        </div>
    </div>
</body>
</html>"""
    return html


def build_job_sheet_excel(job_row: Dict[str, Any], hist_df: pd.DataFrame) -> bytes:
    """Generate an official Excel (.XLSX) Job Sheet workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lembar_Kendali"

    font_title = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    font_sub = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_regular = Font(name="Calibri", size=10)
    font_muted = Font(name="Calibri", size=9, color="64748B")

    fill_primary = PatternFill(start_color="0A3D8F", end_color="0A3D8F", fill_type="solid")
    fill_accent = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_header_tbl = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    col_widths = {"A": 6, "B": 24, "C": 38, "D": 22, "E": 34}
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    ws.merge_cells("A1:E1")
    ws["A1"] = "LEMBAR KENDALI & MONITORING PEKERJAAN"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_primary
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:E2")
    ws["A2"] = "PT PLN INDONESIA POWER SERVICES"
    ws["A2"].font = font_sub
    ws["A2"].fill = fill_primary
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    ws.merge_cells("A4:E4")
    ws["A4"] = "I. INFORMASI IDENTITAS PEKERJAAN & KONTRAK"
    ws["A4"].font = font_section
    ws["A4"].fill = fill_primary
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    info_pairs = [
        ("No. Urut Pekerjaan", job_row.get("NO", "-"), "Status Pekerjaan", job_row.get("STATUS", "-")),
        ("Judul Pekerjaan", job_row.get("Judul Pekerjaan", "-"), "Status Pembayaran", job_row.get("STATUS PEMBAYARAN", "-")),
        ("User / Bidang", job_row.get("User", "-"), "Nama Vendor", job_row.get("Nama Vendor", "-")),
        ("Eksekutor", job_row.get("Eksekutor", "-"), "Nomor Kontrak IPS", job_row.get("Kontrak IPS", "-")),
        ("Proses Kontrak", job_row.get("Proses Kontrak", "-"), "Jadwal Pelaksanaan", f"{job_row.get('Mulai Pekerjaan', '-')} s.d. {job_row.get('Berakhir Pekerjaan', '-')}")
    ]

    curr_row = 5
    for lbl1, val1, lbl2, val2 in info_pairs:
        ws.cell(row=curr_row, column=2, value=str(lbl1)).font = font_bold
        ws.cell(row=curr_row, column=2).fill = fill_accent
        ws.cell(row=curr_row, column=3, value=str(val1)).font = font_regular
        ws.cell(row=curr_row, column=4, value=str(lbl2)).font = font_bold
        ws.cell(row=curr_row, column=4).fill = fill_accent
        ws.cell(row=curr_row, column=5, value=str(val2)).font = font_regular
        for c in range(1, 6):
            ws.cell(row=curr_row, column=c).border = thin_border
        curr_row += 1

    curr_row += 1
    ws.merge_cells(f"A{curr_row}:E{curr_row}")
    ws[f"A{curr_row}"] = "II. RINGKASAN FINANSIAL"
    ws[f"A{curr_row}"].font = font_section
    ws[f"A{curr_row}"].fill = fill_primary
    ws[f"A{curr_row}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    curr_row += 1

    fin_pairs = [
        ("Nilai Kontrak Sebelum PPN (IPS)", _fmt_rp_str(job_row.get("Nilai Kontrak Sebelum PPN (IPS)")), "Nilai Kontrak Mitra (PPN)", _fmt_rp_str(job_row.get("Nilai Kontrak Mitra Setelah PPN"))),
        ("Nilai Kontrak Setelah PPN (IPS)", _fmt_rp_str(job_row.get("Nilai Kontrak Setelah PPN (IPS)")), "Estimasi Margin", _fmt_rp_str(job_row.get("MARGIN"))),
        ("Nilai Invoice Setelah PPN", _fmt_rp_str(job_row.get("Nilai Invoice Setelah PPN")), "Catatan Kendala", str(job_row.get("Kendala", "")).strip() or "-")
    ]
    for lbl1, val1, lbl2, val2 in fin_pairs:
        ws.cell(row=curr_row, column=2, value=str(lbl1)).font = font_bold
        ws.cell(row=curr_row, column=2).fill = fill_accent
        ws.cell(row=curr_row, column=3, value=str(val1)).font = font_regular
        ws.cell(row=curr_row, column=4, value=str(lbl2)).font = font_bold
        ws.cell(row=curr_row, column=4).fill = fill_accent
        ws.cell(row=curr_row, column=5, value=str(val2)).font = font_regular
        for c in range(1, 6):
            ws.cell(row=curr_row, column=c).border = thin_border
        curr_row += 1

    curr_row += 1
    ws.merge_cells(f"A{curr_row}:E{curr_row}")
    ws[f"A{curr_row}"] = "III. RIWAYAT HISTORI & LOG AKTIVITAS"
    ws[f"A{curr_row}"].font = font_section
    ws[f"A{curr_row}"].fill = fill_primary
    ws[f"A{curr_row}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    curr_row += 1

    tbl_headers = ["No", "Tanggal & Waktu", "Aktivitas / Field", "Nilai Lama / Pencatat", "Keterangan / Nilai Baru"]
    for c_idx, h_name in enumerate(tbl_headers, start=1):
        cell = ws.cell(row=curr_row, column=c_idx, value=h_name)
        cell.font = font_title
        cell.fill = fill_header_tbl
        cell.alignment = Alignment(horizontal="center" if c_idx == 1 else "left", vertical="center")
        cell.border = thin_border
    curr_row += 1

    judul_val = str(job_row.get("Judul Pekerjaan", "-"))
    filtered_hist = hist_df[hist_df["Judul Pekerjaan"] == judul_val] if hist_df is not None and not hist_df.empty and "Judul Pekerjaan" in hist_df.columns else pd.DataFrame()

    if not filtered_hist.empty:
        sorted_hist = filtered_hist.sort_values("Timestamp", ascending=False)
        for h_no, (_, h) in enumerate(sorted_hist.iterrows(), 1):
            ws.cell(row=curr_row, column=1, value=h_no).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=2, value=str(h.get("Timestamp", "-")))
            ws.cell(row=curr_row, column=3, value=str(h.get("Field", "-"))).font = font_bold
            ws.cell(row=curr_row, column=4, value=str(h.get("Nilai Lama", "-"))).font = font_muted
            ws.cell(row=curr_row, column=5, value=str(h.get("Nilai Baru", "-")))
            for c in range(1, 6):
                ws.cell(row=curr_row, column=c).border = thin_border
            curr_row += 1
    else:
        ws.merge_cells(f"A{curr_row}:E{curr_row}")
        ws[f"A{curr_row}"] = "Belum ada riwayat histori tercatat untuk pekerjaan ini."
        ws[f"A{curr_row}"].font = font_muted
        ws[f"A{curr_row}"].alignment = Alignment(horizontal="center")
        for c in range(1, 6):
            ws.cell(row=curr_row, column=c).border = thin_border
        curr_row += 1

    curr_row += 2
    ws.cell(row=curr_row, column=2, value="Disiapkan Oleh:").font = font_bold
    ws.cell(row=curr_row, column=5, value="Disetujui / Diverifikasi Oleh:").font = font_bold
    ws.cell(row=curr_row, column=2).alignment = Alignment(horizontal="center")
    ws.cell(row=curr_row, column=5).alignment = Alignment(horizontal="center")

    curr_row += 3
    user_str = str(job_row.get("User", "PIC"))
    ws.cell(row=curr_row, column=2, value=f"( {user_str} )").font = font_bold
    ws.cell(row=curr_row, column=5, value="( ........................................ )").font = font_bold
    ws.cell(row=curr_row, column=2).alignment = Alignment(horizontal="center")
    ws.cell(row=curr_row, column=5).alignment = Alignment(horizontal="center")

    curr_row += 1
    ws.cell(row=curr_row, column=2, value="Pengawas / Pelaksana").font = font_muted
    ws.cell(row=curr_row, column=5, value="Manager / Team Leader").font = font_muted
    ws.cell(row=curr_row, column=2).alignment = Alignment(horizontal="center")
    ws.cell(row=curr_row, column=5).alignment = Alignment(horizontal="center")

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()
